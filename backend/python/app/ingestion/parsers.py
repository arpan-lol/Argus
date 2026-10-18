import csv
import codecs
import gzip
import hashlib
import io
import json
import re
import zipfile
from itertools import chain
from pathlib import Path
from urllib.parse import urlsplit

from defusedxml import ElementTree


SUPPORTED = {".txt", ".md", ".markdown", ".xml", ".jats", ".nxml", ".html", ".htm", ".pdf", ".docx", ".csv", ".tsv", ".xlsx", ".json"}


def _table_text(rows, headers=None):
    normalized = [["" if value is None else str(value).strip() for value in row] for row in rows]
    if not normalized:
        return "", {"rows": [], "headers": headers or []}
    labels = list(headers) if headers is not None else normalized[0]
    data = normalized if headers is not None else normalized[1:]
    rendered = []
    for row_number, row in enumerate(data, 1):
        cells = []
        for index, value in enumerate(row):
            if not value:
                continue
            label = labels[index] if index < len(labels) else f"Column {index + 1}"
            cells.append(f"{label}: {value}" if label else value)
        if cells:
            rendered.append("; ".join(cells))
    structured = {"headers": labels, "rows": normalized, "dataRows": data}
    return "\n".join(rendered) or "\n".join("\t".join(row) for row in normalized), structured


def decode_source_text(raw, partial=False):
    encoding = "utf-16" if raw.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)) else "utf-8-sig"
    try:
        text = codecs.getincrementaldecoder(encoding)(errors="strict").decode(raw, final=not partial)
    except UnicodeError as error:
        raise ValueError("File is not readable UTF-8 or BOM-marked UTF-16 text; no binary-format parser is available") from error
    controls = sum(ord(character) < 32 and character not in "\t\n\r\f" for character in text)
    if "\x00" in text or controls > max(1, len(text) // 200):
        raise ValueError("Binary content cannot be processed by the text reader")
    return text


def parse(path, max_decompressed_bytes=209715200, max_text_chars=20971520, max_archive_members=10000):
    path = Path(path)
    suffix = path.suffix.lower()
    raw = path.read_bytes()
    if suffix == ".gz":
        with gzip.GzipFile(fileobj=io.BytesIO(raw)) as stream:
            raw = stream.read(max_decompressed_bytes + 1)
        if len(raw) > max_decompressed_bytes:
            raise ValueError(f"Decompressed document exceeds configured {max_decompressed_bytes} byte limit")
        suffix = Path(path.stem).suffix.lower()
    if suffix not in SUPPORTED:
        raise ValueError(f"Unsupported format {suffix or '(no extension)'}. Supported: {', '.join(sorted(SUPPORTED))}. OCR and legacy Office binaries are not enabled.")
    if suffix in {".docx", ".xlsx"}:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            members = archive.infolist()
            if len(members) > max_archive_members or sum(member.file_size for member in members) > max_decompressed_bytes:
                raise ValueError("Office archive exceeds configured member or expanded-byte processing limits")
    blocks = []
    text_characters = 0
    metadata = {"filename": path.name, "format": suffix, "parser": "source-blocks-4", "warnings": []}

    def add(text, kind="paragraph", location=None, structured=None):
        nonlocal text_characters
        if text and text.strip():
            text_characters += len(text)
            if text_characters > max_text_chars:
                raise ValueError(f"Parsed text exceeds configured {max_text_chars} character processing limit")
            location = location or {}
            blocks.append({"ordinal": len(blocks), "kind": kind, "role": location.get("role", "table" if kind == "table" else "content"), "text": text, "location": location, "structured": structured})

    if suffix in {".xml", ".jats", ".nxml"}:
        root = ElementTree.fromstring(raw)
        if sum(1 for node in root.iter() if node.tag.split("}")[-1] == "article") > 1:
            raise ValueError("This XML file contains multiple articles. Supply one article per source file so revisions and source locations remain unambiguous.")
        article = root if root.tag.split("}")[-1] == "article" else next((n for n in root.iter() if n.tag.split("}")[-1] == "article"), root)
        metadata["sourceDocumentType"] = article.get("article-type", "unspecified")
        body = next((node for node in article if node.tag.split("}")[-1] == "body"), None)
        metadata["authors"] = []
        for node in article.iter():
            tag = node.tag.split("}")[-1]
            if tag == "license":
                metadata["license"] = " ".join("".join(node.itertext()).split())
                metadata["licenseUrl"] = next((value for key, value in node.attrib.items() if key.endswith("href")), None)
            elif tag == "contrib" and node.get("contrib-type", "author") == "author":
                metadata["authors"].append(" ".join(" ".join(node.itertext()).split()))
            elif tag == "pub-date":
                values = {part.tag.split("}")[-1]: "".join(part.itertext()).strip() for part in node}
                metadata.setdefault("publicationDates", []).append({"type": node.get("pub-type") or node.get("publication-format"), **values})
                if "year" in values:
                    metadata.setdefault("year", values["year"])
        title = next((n for n in article.iter() if n.tag.split("}")[-1] == "article-title"), None)
        if title is not None:
            metadata["title"] = "".join(title.itertext()).strip()
        for node in article.iter():
            tag = node.tag.split("}")[-1]
            if tag == "article-id" and node.get("pub-id-type") in {"pmc", "pmcid", "doi"}:
                metadata[node.get("pub-id-type")] = "".join(node.itertext()).strip()
        if metadata.get("pmc") or metadata.get("pmcid"):
            ident = metadata.get("pmcid") or metadata["pmc"]
            metadata["sourceUrl"] = f"https://pmc.ncbi.nlm.nih.gov/articles/{ident if ident.startswith('PMC') else 'PMC' + ident}/"
        containers = {"p", "title", "article-title", "table", "disp-formula", "list-item", "caption", "mixed-citation", "element-citation", "preformat"}
        roles = {"front": "front_matter", "body": "body", "back": "back_matter", "abstract": "abstract", "trans-abstract": "abstract", "ref-list": "references", "ref": "references", "table": "table", "caption": "caption"}

        def non_whitespace(text):
            return "".join(text.split())

        def walk(node, ancestry, role="content", region="content", section_path=()):
            tag = node.tag.split("}")[-1]
            structural = f"{ancestry}/{tag}"
            role = roles.get(tag, role)
            if tag in {"front", "body", "back"}:
                region = tag
            if tag == "sec":
                heading = next((child for child in node if child.tag.split("}")[-1] == "title"), None)
                section_path = (*section_path, {"title": "".join(heading.itertext()).strip() if heading is not None else "", "structuralPath": structural, "sourceId": node.get("id")})
            location = {"structuralPath": structural, "element": tag, "role": role, "sourceRegion": region, "sectionPath": list(section_path)}
            if node.get("id"):
                location["sourceId"] = node.get("id")
            nested_blocks = any(child.tag.split("}")[-1] in containers for child in node.iter() if child is not node)
            if tag == "table" or (tag in containers and not nested_blocks):
                if tag == "table":
                    rows = [["".join(cell.itertext()).strip() for cell in row if cell.tag.split("}")[-1] in {"td", "th"}] for row in node.iter() if row.tag.split("}")[-1] == "tr"]
                    source_text = "".join(node.itertext())
                    row_text, table_structured = _table_text(rows)
                    retained = row_text if non_whitespace(row_text) == non_whitespace(source_text) else source_text
                    add(retained, "table", location, {**table_structured, "textRendering": "rows_with_headers" if retained == row_text else "source_text"})
                else:
                    kind = "citation" if tag in {"mixed-citation", "element-citation"} else "heading" if tag in {"title", "article-title"} else "preformatted" if tag == "preformat" else "caption" if tag == "caption" else "paragraph"
                    add("".join(node.itertext()), kind, location)
                return
            kind = "citation" if role == "references" else "caption" if role == "caption" else "paragraph"
            add(node.text, kind, {**location, "textPosition": "text"})
            for ordinal, child in enumerate(node):
                walk(child, f"{structural}[{ordinal}]", role, region, section_path)
                add(child.tail, kind, {**location, "textPosition": "tail", "afterChild": ordinal})
        walk(article, "")
        body_text = "".join(body.itertext()) if body is not None else ""
        body_available = non_whitespace(body_text)
        body_retained = non_whitespace("".join(block["text"] for block in blocks if block["location"].get("sourceRegion") == "body"))
        complete = body is not None and body_available == body_retained
        reasons = []
        if body is None:
            reasons.append("No JATS body is present; available front matter, abstract and back matter do not establish complete article body coverage.")
        elif not body_available:
            reasons.append("The JATS body contains no readable text; images or external content may require a different reader.")
        elif not complete:
            reasons.append("Retained body text differs from the original ordered body text; processing requires attention before extraction.")
        metadata["sourceCoverage"] = {
            "bodyPresent": body is not None,
            "bodyTextCharacters": len(body_text.strip()),
            "bodyAvailableNonWhitespaceCharacters": len(body_available),
            "bodyRetainedNonWhitespaceCharacters": len(body_retained),
            "bodyMissingNonWhitespaceCharacters": max(0, len(body_available) - len(body_retained)),
            "bodyCoverageComplete": complete and bool(body_available),
            "bodyAvailableTextSha256": hashlib.sha256(body_available.encode("utf-8")).hexdigest(),
            "bodyRetainedTextSha256": hashlib.sha256(body_retained.encode("utf-8")).hexdigest(),
            "accounting": "exact_ordered_non_whitespace_characters",
            "scope": "available_jats_text",
            "requiresAttention": bool(body_available) and not complete,
            "attentionReasons": reasons,
        }
        metadata["warnings"].extend(reasons)
    elif suffix == ".pdf":
        from pypdf import PdfReader
        page_count = 0
        readable_pages = 0
        for index, page in enumerate(PdfReader(io.BytesIO(raw)).pages):
            page_count += 1
            text = page.extract_text() or ""
            readable_pages += bool(text.strip())
            add(text, location={"pageNumber": index + 1})
            if not text.strip():
                metadata["warnings"].append(f"Page {index + 1} has no extractable text; OCR is not configured")
        metadata["warnings"].append("PDF text reading order is parser-derived; coordinates and table structure are not reconstructed")
        if readable_pages < page_count:
            metadata["sourceCoverage"] = {"scope": "available_pdf_text", "pageCount": page_count, "pagesWithText": readable_pages, "pagesWithoutText": [i + 1 for i, page in enumerate(PdfReader(io.BytesIO(raw)).pages) if not (page.extract_text() or "").strip()], "coverageLabel": "partial_pdf_text", "limitedSource": True, "requiresAttention": True, "attentionReasons": ["One or more pages have no extractable text; OCR is unavailable, so processing must not continue as complete source coverage."]}
    elif suffix == ".docx":
        from docx import Document
        document = Document(io.BytesIO(raw))
        from docx.oxml.ns import qn
        from docx.text.paragraph import Paragraph
        from docx.table import Table
        section_path = []
        for index, element in enumerate(document.element.body):
            if element.tag == qn("w:p"):
                paragraph = Paragraph(element, document)
                is_heading = paragraph.style.name.startswith("Heading")
                if is_heading:
                    level = int(re.search(r"\d+", paragraph.style.name).group()) if re.search(r"\d+", paragraph.style.name) else 1
                    section_path = section_path[:max(0, level - 1)] + [paragraph.text.strip()]
                add(paragraph.text, "heading" if is_heading else "paragraph", {"structuralPath": f"body/{index}", "sectionPath": list(section_path)})
            elif element.tag == qn("w:tbl"):
                rows = [[cell.text for cell in row.cells] for row in Table(element, document).rows]
                rendered, structured = _table_text(rows)
                add(rendered, "table", {"structuralPath": f"body/{index}", "sectionPath": list(section_path)}, structured)
    elif suffix == ".xlsx":
        from openpyxl import load_workbook
        from openpyxl.utils import get_column_letter
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
        cached = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        metadata['warnings'].append('Spreadsheet formulas are retained and never executed; cached values may be absent or stale')
        metadata['sheetStates'] = {sheet.title: sheet.sheet_state for sheet in workbook}
        for sheet in workbook:
            rows = iter(sheet.iter_rows())
            first = next(rows, None)
            if first is None:
                continue
            headers = [str(cell.value) if cell.value is not None else f"Column {column}" for column, cell in enumerate(first, 1)]
            cached_rows = iter(cached[sheet.title].iter_rows())
            for index, row in enumerate(chain([first], rows), start=1):
                cells = ["" if cell.value is None else str(cell.value) for cell in row]
                rendered, structured = ("\t".join(cells), {"headers": headers, "rows": [cells], "isHeader": True}) if index == 1 else _table_text([cells], headers=headers)
                cached_row = next(cached_rows)
                structured["cells"] = [{"coordinate": f"{get_column_letter(column)}{index}", "value": None if cell.value is None else str(cell.value), "formula": str(cell.value) if cell.data_type == "f" else None, "cachedValue": None if previous.value is None else str(previous.value), "dataType": cell.data_type} for column, (cell, previous) in enumerate(zip(row, cached_row, strict=True), 1)]
                add(rendered, "table", {"sheetName": sheet.title, "rowNumber": index, "cellRange": f"A{index}:{get_column_letter(len(row))}{index}"}, structured)
        workbook.close()
        cached.close()
    elif suffix in {".csv", ".tsv"}:
        rows = list(csv.reader(io.StringIO(decode_source_text(raw)), delimiter="\t" if suffix == ".tsv" else ","))
        headers = rows[0] if rows else []
        for index, row in enumerate(rows):
            rendered, structured = _table_text([row], headers=headers if index else None)
            structured.update({"cells": row, "rowNumber": index + 1, "isHeader": index == 0})
            add(rendered, "table", {"structuralPath": f"row/{index + 1}", "rowNumber": index + 1}, structured)
    elif suffix in {".html", ".htm"}:
        from bs4 import BeautifulSoup
        soup = BeautifulSoup(raw, "html.parser")
        for element in soup(["script", "style", "nav"]):
            element.decompose()
        metadata["title"] = soup.title.get_text() if soup.title else path.stem
        section_path = []
        for index, element in enumerate(soup.find_all(["h1", "h2", "h3", "h4", "h5", "h6", "p", "pre", "table", "li"])):
            if element.find_parent(["table", "li", "pre"]):
                continue
            is_heading = bool(re.fullmatch("h[1-6]", element.name))
            if is_heading:
                level = int(element.name[1])
                section_path = section_path[:level - 1] + [element.get_text(" ", strip=True)]
            location = {"structuralPath": f"{element.name}/{index}", "sectionPath": list(section_path)}
            if element.name == "table":
                rows = [[cell.get_text(" ", strip=True) for cell in row.find_all(["th", "td"], recursive=False)] for row in element.find_all("tr")]
                rendered, structured = _table_text(rows)
                add(rendered, "table", location, structured)
            else:
                add(element.get_text(" ", strip=True), "heading" if is_heading else "preformatted" if element.name == "pre" else "list_item" if element.name == "li" else "paragraph", location)
    elif suffix in {".md", ".markdown"}:
        text = decode_source_text(raw)
        lines = text.splitlines()
        section_path = []
        index = 0
        while index < len(lines):
            line = lines[index]
            stripped = line.strip()
            if not stripped:
                index += 1
                continue
            heading = re.match(r"^(#{1,6})\s+(.+?)\s*#*\s*$", stripped)
            if heading:
                level, value = len(heading.group(1)), heading.group(2).strip()
                section_path = section_path[:level - 1] + [value]
                add(value, "heading", {"structuralPath": f"line/{index + 1}", "lineNumber": index + 1, "headingLevel": level, "sectionPath": list(section_path)})
                index += 1
                continue
            if "|" in stripped and index + 1 < len(lines) and re.match(r"^\s*\|?\s*:?-{3,}", lines[index + 1]):
                table_lines = [stripped]
                index += 2
                while index < len(lines) and "|" in lines[index] and lines[index].strip():
                    table_lines.append(lines[index].strip())
                    index += 1
                rows = [[cell.strip() for cell in row.strip("|").split("|")] for row in [table_lines[0], *table_lines[1:]]]
                rendered, structured = _table_text(rows)
                add(rendered, "table", {"structuralPath": f"line/{index - len(table_lines) + 1}", "sectionPath": list(section_path)}, structured)
                continue
            start = index
            paragraph = [stripped]
            index += 1
            while index < len(lines) and lines[index].strip() and not re.match(r"^#{1,6}\s+", lines[index].strip()):
                paragraph.append(lines[index].strip())
                index += 1
            add(" ".join(paragraph), "paragraph", {"structuralPath": f"line/{start + 1}", "lineNumber": start + 1, "sectionPath": list(section_path)})
    else:
        text = decode_source_text(raw)
        if suffix == ".json":
            value = json.loads(text)
            records = value if isinstance(value, list) else value.get("records", []) if isinstance(value, dict) and isinstance(value.get("records"), list) else [value]
            for index, record in enumerate(records):
                rendered = json.dumps(record, ensure_ascii=False, sort_keys=True)
                add(rendered, "record", {"structuralPath": f"record/{index + 1}", "recordNumber": index + 1}, {"record": record})
        else:
            front = text[:4096].split('\n\n', 1)[0]
            for line in front.splitlines():
                key, separator, value = line.partition(':')
                value = value.strip()
                if separator and key.strip().casefold() == 'title' and value:
                    metadata['title'] = value[:500]
                if separator and key.strip().casefold() in {'original url', 'source url'}:
                    parsed_url = urlsplit(value)
                    if parsed_url.scheme in {'http', 'https'} and parsed_url.hostname and not parsed_url.username and not parsed_url.password:
                        metadata['sourceUrl'] = value
            for index, paragraph in enumerate(re.split(r"\n\s*\n", text)):
                add(paragraph, "paragraph", {"structuralPath": f"block/{index}"})
    if not blocks:
        raise ValueError("Parser produced no readable text. This document may require OCR or a different parser.")
    return {"blocks": blocks, "metadata": metadata}


def normalize(parsed):
    blocks = []
    parts = []
    cursor = 0
    for block in parsed["blocks"]:
        text = block["text"].replace("\r\n", "\n").replace("\r", "\n").strip()
        if not text:
            continue
        if parts:
            cursor += 2
        blocks.append({**block, "text": text, "startOffset": cursor, "endOffset": cursor + len(text), "offsetUnit": "unicode_code_point", "fidelity": "derived"})
        parts.append(text)
        cursor += len(text)
    return {"text": "\n\n".join(parts), "blocks": blocks, "metadata": parsed["metadata"], "normalization": "Newlines normalized; block edges stripped; two-newline block separator. Locations refer to original structural/page positions; offsets refer to normalized text, not raw bytes."}
